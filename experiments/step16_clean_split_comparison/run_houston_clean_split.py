# -*- coding: utf-8 -*-
"""Fair Houston2018 comparison with a train-only 90/10 development split.

【中文说明：这是目前三个数据集共用的正式训练/测试入口】
--method baseline      ：原始三模态MSAF3；
--method dynamic       ：Step15，Transformer1输入动态加权（旧版/消融）；
--method dynamic_final ：Step17，最终三分支求和前动态加权（当前最终版）。

本文件负责数据划分、300轮训练、验证集选最优checkpoint、断点续训、
调用全样本测试器以及生成单次实验的TXT/Excel。具体模型结构分别放在
msaf3、Step15和Step17中，本文件主要是“实验流程控制器”。

The official test split is never used during training or checkpoint selection.
Both the original three-modal MSAF3 baseline and the complete dynamic method
use the exact same stratified split.  After training, the selected checkpoint
is evaluated on all 2,000,160 official test pixels and an MSAF-style text
report plus one Excel workbook are written for that run.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import random
import subprocess
import sys
import time
from pathlib import Path

import numpy as np
import torch
import torch.nn as nn
from openpyxl import Workbook
from openpyxl.styles import Font
from sklearn.metrics import confusion_matrix
from torch.utils.data import DataLoader


HERE = Path(__file__).resolve().parent
EXPERIMENTS = HERE.parent
ROOT = EXPERIMENTS.parent
MSAF3_PATH = ROOT / "msaf3" / "train_msaf3.py"
STEP15_PATH = EXPERIMENTS / "step15_warmup_dynamic_300" / "train_warmup_then_dynamic.py"
STEP17_PATH = EXPERIMENTS / "step17_final_fusion_dynamic" / "train_final_fusion_dynamic.py"
FULL_TEST_PATH = EXPERIMENTS / "step14_official_full_test" / "run_official_full_test.py"


def load_module(name: str, path: Path):
    """按路径加载不同实验模型，避免复制基础网络代码。"""
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


MSAF3 = load_module("clean_split_msaf3", MSAF3_PATH)
DYNAMIC = load_module("clean_split_dynamic", STEP15_PATH)
DYNAMIC_FINAL = load_module("clean_split_dynamic_final", STEP17_PATH)
BASE = DYNAMIC.BASE
DEVICE = BASE.DEVICE


def is_dynamic_method(method: str) -> bool:
    """判断是否属于带warm-up和动态权重的两种方法。"""
    return method in ("dynamic", "dynamic_final")


def parse_args() -> argparse.Namespace:
    """解析数据集、方法、训练参数、损失权重和FW超参数。"""
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--method",
        choices=("baseline", "dynamic", "dynamic_final"),
        required=True,
    )
    parser.add_argument(
        "--dataset",
        choices=("houston2018", "szutree-r1", "szutree-r2"),
        default="houston2018",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=None,
    )
    parser.add_argument("--output-root", type=Path, default=HERE / "results")
    parser.add_argument("--seed", type=int, default=4)
    parser.add_argument("--val-ratio", type=float, default=0.10)
    parser.add_argument("--train-percent-per-class", type=float, default=1.0)
    parser.add_argument("--epochs", type=int, default=300)
    parser.add_argument("--warmup-epochs", type=int, default=50)
    parser.add_argument("--batch-size", type=int, default=16)
    parser.add_argument("--eval-batch-size", type=int, default=256)
    parser.add_argument("--patch-size", type=int, default=11)
    parser.add_argument("--learning-rate", type=float, default=2e-4)
    parser.add_argument("--weight-decay", type=float, default=5e-3)
    parser.add_argument("--aux-weight", type=float, default=0.1)
    parser.add_argument("--formula16-weight", type=float, default=0.2)
    parser.add_argument("--q-loss-weight", type=float, default=0.1)
    parser.add_argument("--fw-alpha", type=float, default=1.0)
    parser.add_argument("--fw-beta", type=float, default=0.25)
    parser.add_argument("--fw-gamma", type=float, default=64.0)
    parser.add_argument("--fw-iterations", type=int, default=10)
    parser.add_argument("--q-balance-lambda", type=float, default=0.05)
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--pin-memory", action="store_true")
    parser.add_argument("--no-pretrained", action="store_true")
    parser.add_argument("--skip-full-test", action="store_true")
    parser.add_argument("--smoke-epochs", type=int, default=0, help=argparse.SUPPRESS)
    return parser.parse_args()


def set_seed(seed: int) -> None:
    """固定随机种子，保证baseline与改进方法的数据划分和初始化可复现。"""
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed(seed)
        torch.cuda.manual_seed_all(seed)


def stratified_train_validation(
    official_train: np.ndarray, val_ratio: float, seed: int
) -> tuple[np.ndarray, np.ndarray]:
    """从Houston官方训练集内部按类别划分训练/验证集。

    每一类约90%训练、10%验证；官方测试集完全不参与checkpoint选择。
    """
    if not 0.0 < val_ratio < 1.0:
        raise ValueError("--val-ratio must be between 0 and 1")
    rng = np.random.RandomState(seed)
    train_parts, val_parts = [], []
    for label in np.unique(official_train[:, 0]):
        class_rows = official_train[official_train[:, 0] == label]
        order = rng.permutation(len(class_rows))
        val_count = max(1, int(round(len(class_rows) * val_ratio)))
        val_parts.append(class_rows[order[:val_count]])
        train_parts.append(class_rows[order[val_count:]])
    train_idx = np.concatenate(train_parts, axis=0)
    val_idx = np.concatenate(val_parts, axis=0)
    return train_idx[rng.permutation(len(train_idx))], val_idx[rng.permutation(len(val_idx))]


def split_hash(train_idx: np.ndarray, val_idx: np.ndarray) -> str:
    """计算划分指纹；baseline与改进方法哈希相同才是公平对比。"""
    digest = hashlib.sha256()
    digest.update(np.ascontiguousarray(train_idx).tobytes())
    digest.update(np.ascontiguousarray(val_idx).tobytes())
    return digest.hexdigest()


def counts(index_array: np.ndarray, label_shift: int, classes: int) -> np.ndarray:
    """统计索引数组中每个类别的样本数。"""
    return np.bincount(
        index_array[:, 0].astype(np.int64) - label_shift, minlength=classes
    )


def loader_kwargs(args: argparse.Namespace, shuffle: bool) -> dict:
    """统一构造DataLoader参数；训练集shuffle，验证集不shuffle。"""
    result = {
        "shuffle": shuffle,
        "num_workers": args.num_workers,
        "pin_memory": args.pin_memory,
    }
    if args.num_workers > 0:
        result["persistent_workers"] = True
    return result


def fusion_logits(output):
    """兼容baseline直接返回logits和动态模型返回字典两种接口。"""
    return output["logits_fusion"] if isinstance(output, dict) else output


@torch.no_grad()
def evaluate(model, loader, criterion) -> dict[str, float]:
    """在验证集计算融合损失、OA及权重集中程度。

    mean_max_weight越接近1，表示权重越容易集中到单个模态，可用于监控
    FW权重塌缩。验证阶段不更新任何参数。
    """
    saved_mask = BASE.temporarily_disable_attention_mask(model)
    model.eval()
    total = correct = batches = 0
    loss_sum = max_weight_sum = 0.0
    try:
        for hsi, lidar, rgb, target in loader:
            hsi = hsi.to(DEVICE, non_blocking=True)
            lidar = lidar.to(DEVICE, non_blocking=True)
            rgb = rgb.to(DEVICE, non_blocking=True)
            target = target.to(DEVICE, non_blocking=True)
            output = model(hsi, lidar, rgb)
            logits = fusion_logits(output)
            loss_sum += float(criterion(logits, target).item())
            correct += int(logits.argmax(dim=1).eq(target).sum().item())
            total += int(target.numel())
            if isinstance(output, dict) and "weights" in output:
                max_weight_sum += float(output["weights"].max(dim=1).values.mean().item())
            batches += 1
    finally:
        BASE.restore_attention_mask(saved_mask)
    return {
        "loss": loss_sum / max(1, batches),
        "OA": correct / max(1, total),
        "mean_max_weight": max_weight_sum / max(1, batches) if batches else 0.0,
    }


def train_epoch(model, loader, optimizer, criterion, args) -> dict[str, float]:
    """训练一轮。

    baseline只优化最终分类CE；两种动态方法共同使用Step15定义的完整损失，
    即融合CE、三个辅助CE、高置信错误校准以及q正确性BCE。
    """
    model.train()
    total = correct = batches = 0
    sums = np.zeros(5, dtype=np.float64)
    for hsi, lidar, rgb, target in loader:
        hsi = hsi.to(DEVICE, non_blocking=True)
        lidar = lidar.to(DEVICE, non_blocking=True)
        rgb = rgb.to(DEVICE, non_blocking=True)
        target = target.to(DEVICE, non_blocking=True)
        optimizer.zero_grad()
        output = model(hsi, lidar, rgb)
        if is_dynamic_method(args.method):
            losses = DYNAMIC.compute_loss(output, target, args, criterion)
            total_loss, fusion_loss, aux_loss, calibration_loss, q_loss = losses
        else:
            fusion_loss = criterion(output, target)
            total_loss = fusion_loss
            aux_loss = calibration_loss = q_loss = torch.zeros_like(total_loss)
        total_loss.backward()
        optimizer.step()
        logits = fusion_logits(output)
        correct += int(logits.argmax(dim=1).eq(target).sum().item())
        total += int(target.numel())
        sums += [
            float(total_loss.item()),
            float(fusion_loss.item()),
            float(aux_loss.item()),
            float(calibration_loss.item()),
            float(q_loss.item()),
        ]
        batches += 1
    sums /= max(1, batches)
    return {
        "train_total_loss": sums[0],
        "train_fusion_loss": sums[1],
        "train_aux_loss": sums[2],
        "train_formula16_loss": sums[3],
        "train_q_loss": sums[4],
        "train_OA": correct / max(1, total),
    }


def checkpoint_payload(model, optimizer, scheduler, epoch, best_val, best_epoch, args):
    """构造断点续训包，保存模型、优化器、调度器、轮次和完整配置。"""
    return {
        "model_state_dict": {
            key: value.detach().cpu().clone() for key, value in model.state_dict().items()
        },
        "optimizer_state_dict": optimizer.state_dict(),
        "scheduler_state_dict": scheduler.state_dict(),
        "epoch": int(epoch),
        "phase": args.method,
        "best_val_OA": float(best_val),
        "best_epoch": int(best_epoch),
        "config": vars(args),
    }


def atomic_save(payload: dict, destination: Path) -> None:
    """安全保存checkpoint：临时文件写完后再替换，降低中断损坏风险。"""
    temporary = destination.with_name(f".{destination.name}.{time.time_ns()}.tmp")
    torch.save(payload, temporary)
    for attempt in range(10):
        try:
            temporary.replace(destination)
            return
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(1.0)


def append_epoch_log(path: Path, row: dict) -> None:
    """逐轮追加训练/验证指标；汇总Excel的曲线和时间来自此文件。"""
    exists = path.exists()
    with path.open("a", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row.keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(row)


def compute_macs_params(model, dataset, patch_size: int) -> tuple[float, float]:
    """统计参数量，并尽可能使用thop估算单样本MACs。"""
    params = float(sum(parameter.numel() for parameter in model.parameters()))
    try:
        from thop import profile

        hsi = torch.randn(1, dataset["hsi_channels"], patch_size, patch_size, device=DEVICE)
        lidar = torch.randn(1, dataset["lidar_channels"], patch_size, patch_size, device=DEVICE)
        rgb = torch.randn(1, dataset["rgb_channels"], patch_size, patch_size, device=DEVICE)
        saved_mask = BASE.temporarily_disable_attention_mask(model)
        model.eval()
        try:
            macs, _ = profile(model, inputs=(hsi, lidar, rgb), verbose=False)
        finally:
            BASE.restore_attention_mask(saved_mask)
        return float(macs), params
    except Exception as error:
        print(f"[warn] MAC profiling failed: {error}", flush=True)
        return 0.0, params


def write_msaf_report(
    path: Path,
    train_counts: np.ndarray,
    val_counts: np.ndarray,
    metrics: dict,
    training_seconds: float,
    testing_seconds: float,
    params: float,
    macs: float,
) -> None:
    """写出与原MSAF报告风格一致的文本结果。"""
    oa, aa, kappa = metrics["OA"], metrics["AA"], metrics["kappa"]
    each = np.asarray(metrics["per_class_accuracy"], dtype=np.float64)
    values = each.tolist() + [oa, aa, kappa]
    lines = [
        f"training samples are:{tuple(int(x) for x in train_counts)}",
        f"validation samples are:{tuple(int(x) for x in val_counts)}",
        f"OAs for each iteration are:{[oa]}",
        f"AAs for each iteration are:{[aa]}",
        f"KAPPAs for each iteration are:{[kappa]}",
        "",
        f"mean_OA ± std_OA is: {oa} ± 0.0",
        f"mean_AA ± std_AA is: {aa} ± 0.0",
        f"mean_KAPPA ± std_KAPPA is: {kappa} ± 0.0",
        "",
        f"Mean of all elements in confusion matrix: {each}",
        f"Standard deviation of all elements in confusion matrix: {np.zeros_like(each)}",
        "",
        "Per-class accuracy (%): "
        + " ".join(f"C{i}:{accuracy * 100:.2f}" for i, accuracy in enumerate(each)),
        "Per-class accuracy with std (%): "
        + " ".join(f"C{i}:{accuracy * 100:.2f}±0.00" for i, accuracy in enumerate(each)),
        "",
        f"All values without std: {values}",
        "",
        "All values with std: " + ", ".join(f"{value} ± 0.0" for value in values) + ", ",
        "",
        f"train time for each iteration are:{[training_seconds]}",
        "",
        f"test time for each iteration are:{[testing_seconds]}",
        "",
        f"all iters acc:{each[None, :]}",
        "",
        f"The models parameters and Macs are:{params} {macs}",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_excel(
    path: Path,
    args,
    split_info: dict,
    training_summary: dict,
    metrics: dict,
    epoch_log: Path,
) -> None:
    """为当前单次实验生成Excel：总体、逐类、混淆矩阵、曲线和配置。"""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    fields = {
        "Dataset": args.dataset,
        "Method": args.method,
        "Seed": args.seed,
        "Validation ratio": args.val_ratio,
        "Split SHA256": split_info["split_sha256"],
        "Best epoch": training_summary["best_epoch"],
        "Best validation OA": training_summary["best_val_OA"],
        "Official test samples": metrics["samples"],
        "Test OA": metrics["OA"],
        "Test AA": metrics["AA"],
        "Test Kappa": metrics["kappa"],
        "Training seconds": training_summary["training_seconds"],
        "Testing seconds": metrics["elapsed_seconds"],
        "Parameters": training_summary["parameters"],
        "MACs": training_summary["macs"],
    }
    for key, value in fields.items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 72

    per_class = workbook.create_sheet("PerClass")
    per_class.append(["Class", "Train", "Validation", "Official test", "Accuracy"])
    for cell in per_class[1]:
        cell.font = Font(bold=True)
    for index, accuracy in enumerate(metrics["per_class_accuracy"]):
        per_class.append(
            [
                index,
                split_info["train_counts"][index],
                split_info["validation_counts"][index],
                metrics["per_class_count"][index],
                accuracy,
            ]
        )

    matrix = workbook.create_sheet("ConfusionMatrix")
    matrix.append(["True\\Pred"] + list(range(len(metrics["confusion_matrix"]))))
    for cell in matrix[1]:
        cell.font = Font(bold=True)
    for index, row in enumerate(metrics["confusion_matrix"]):
        matrix.append([index] + row)

    epochs = workbook.create_sheet("EpochLog")
    with epoch_log.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            epochs.append(row)
    for cell in epochs[1]:
        cell.font = Font(bold=True)

    config = workbook.create_sheet("Config")
    config.append(["Argument", "Value"])
    config["A1"].font = config["B1"].font = Font(bold=True)
    for key, value in sorted(vars(args).items()):
        config.append([key, str(value)])
    workbook.save(path)


def main() -> None:
    """依次完成：准备划分 -> 训练/续训 -> 选最优 -> 全测试 -> 报告。"""
    args = parse_args()
    if args.smoke_epochs > 0:
        # 仅供快速代码连通性测试；正式命令不传此隐藏参数。
        args.epochs = args.smoke_epochs
        args.skip_full_test = True
    default_data_dirs = {
        "houston2018": Path(r"D:\DATA_3\Houston2018\prepared_msaf3"),
        "szutree-r1": Path(r"D:\DATA_3\SZUTreeData2.0\SZUTreeData_R1_2.0\prepared_msaf3"),
        "szutree-r2": Path(r"D:\DATA_3\SZUTreeData2.0\SZUTreeData_R2_2.0\prepared_msaf3"),
    }
    if args.data_dir is None:
        # 三个数据集各自使用已经预处理好的缓存目录。
        args.data_dir = default_data_dirs[args.dataset]
    dataset_tag = args.dataset.replace("-", "_")
    method_dir = args.output_root / f"{dataset_tag}_{args.method}_seed{args.seed}"
    # 方法名进入目录名，因此baseline、dynamic、dynamic_final不会互相覆盖。
    method_dir.mkdir(parents=True, exist_ok=True)
    set_seed(args.seed)

    dataset = MSAF3.load_cache(args.data_dir)
    num_classes = int(dataset["num_classes"])
    if args.dataset == "houston2018":
        # Houston：官方train内部90/10；官方test保持原样。
        official_train = np.asarray(dataset["train_idx"], dtype=np.int32)
        train_idx, val_idx = stratified_train_validation(
            official_train, args.val_ratio, args.seed
        )
        test_idx = np.asarray(dataset["test_idx"], dtype=np.int32)
        label_shift = int(dataset["label_shift"])
        source_counts = counts(official_train, label_shift, num_classes)
        source_samples = len(official_train) + len(test_idx)
        development_samples = len(official_train)
    else:
        # SZUTree：先按类别抽1% development，再从其中抽10%作validation；
        # 即约0.9%训练、0.1%验证、其余约99%测试。
        development_idx, test_idx = MSAF3.split_per_class_percent(
            dataset["labels"], args.seed, args.train_percent_per_class
        )
        train_idx, val_idx = MSAF3.split_index_train_validation(
            development_idx, args.seed + 10000, args.val_ratio
        )
        label_shift = int(
            min(train_idx[:, 0].min(), val_idx[:, 0].min(), test_idx[:, 0].min())
        )
        source_counts = counts(
            np.concatenate([development_idx, test_idx], axis=0),
            label_shift,
            num_classes,
        )
        source_samples = len(development_idx) + len(test_idx)
        development_samples = len(development_idx)
    train_counts = counts(train_idx, label_shift, num_classes)
    val_counts = counts(val_idx, label_shift, num_classes)
    if args.dataset == "houston2018" and not np.array_equal(
        train_counts + val_counts, source_counts
    ):
        raise RuntimeError("The 90/10 split does not reconstruct official train_idx")
    if args.dataset != "houston2018" and len(train_idx) + len(val_idx) != development_samples:
        raise RuntimeError("The inner SZUTree train/validation split is incomplete")
    split_digest = split_hash(train_idx, val_idx)
    split_info = {
        "seed": args.seed,
        "validation_ratio": args.val_ratio,
        "split_sha256": split_digest,
        "dataset": args.dataset,
        "source_labeled_samples": int(source_samples),
        "development_samples_before_validation": int(development_samples),
        "train_samples": int(len(train_idx)),
        "validation_samples": int(len(val_idx)),
        "test_samples": int(len(test_idx)),
        "train_counts": train_counts.tolist(),
        "validation_counts": val_counts.tolist(),
        "source_class_counts": source_counts.tolist(),
    }
    np.save(method_dir / "train_idx.npy", train_idx)
    np.save(method_dir / "validation_idx.npy", val_idx)
    np.save(method_dir / "test_idx.npy", test_idx)
    (method_dir / "split.json").write_text(
        json.dumps(split_info, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(split_info, ensure_ascii=False, indent=2), flush=True)

    train_dataset = MSAF3.PreparedTriModalDataset(
        dataset["features"], dataset["rgb"], train_idx, label_shift,
        args.patch_size, dataset["hsi_channels"], dataset["lidar_channels"],
    )
    val_dataset = MSAF3.PreparedTriModalDataset(
        dataset["features"], dataset["rgb"], val_idx, label_shift,
        args.patch_size, dataset["hsi_channels"], dataset["lidar_channels"],
    )
    train_loader = DataLoader(
        train_dataset, batch_size=args.batch_size, **loader_kwargs(args, True)
    )
    val_loader = DataLoader(
        val_dataset, batch_size=args.batch_size, **loader_kwargs(args, False)
    )

    if args.method == "baseline":
        # 原始三模态MSAF3，不含单模态辅助头、q和FW动态权重。
        model = MSAF3.MCF3(
            HSIband=dataset["hsi_channels"],
            lidarband=dataset["lidar_channels"],
            rgbband=dataset["rgb_channels"],
            num_classes=num_classes,
            use_pretrained=not args.no_pretrained,
            use_rgb_pretrained=not args.no_pretrained,
        ).to(DEVICE)
    elif args.method == "dynamic":
        # Step15旧版：权重乘在Transformer1输入上，用于权重位置消融。
        model = DYNAMIC.WarmupDynamicMSAF3(
            dataset["hsi_channels"], dataset["lidar_channels"],
            dataset["rgb_channels"], num_classes, args,
        ).to(DEVICE)
        model.dynamic_enabled = False
    else:
        # Step17最终版：两个Transformer保持原流程，仅最终求和动态加权。
        model = DYNAMIC_FINAL.FinalFusionDynamicMSAF3(
            dataset["hsi_channels"], dataset["lidar_channels"],
            dataset["rgb_channels"], num_classes, args,
        ).to(DEVICE)
        model.dynamic_enabled = False

    optimizer = torch.optim.AdamW(
        model.parameters(), lr=args.learning_rate, weight_decay=args.weight_decay
    )
    scheduler = torch.optim.lr_scheduler.StepLR(optimizer, step_size=50, gamma=0.7)
    criterion = nn.CrossEntropyLoss()
    latest_path = method_dir / "latest.pt"
    best_path = method_dir / "best_model.pt"
    epoch_log = method_dir / "epoch_log.csv"
    best_val, best_epoch, start_epoch = -1.0, 0, 1
    training_seconds = 0.0

    if latest_path.exists():
        # latest.pt用于续训：从保存epoch的下一轮恢复模型、优化器和学习率。
        # best_model.pt只保存验证OA最优状态，最后全测试使用它。
        state = torch.load(latest_path, map_location="cpu", weights_only=False)
        model.load_state_dict(state["model_state_dict"], strict=True)
        optimizer.load_state_dict(state["optimizer_state_dict"])
        scheduler.load_state_dict(state["scheduler_state_dict"])
        start_epoch = int(state["epoch"]) + 1
        best_val = float(state["best_val_OA"])
        best_epoch = int(state["best_epoch"])
        training_seconds = float(state.get("training_seconds", 0.0))
        print(f"Resuming {args.method} at epoch {start_epoch}", flush=True)

    for epoch in range(start_epoch, args.epochs + 1):
        if is_dynamic_method(args.method):
            # 前warmup_epochs轮固定1/3；之后每个样本启用FW动态权重。
            model.dynamic_enabled = epoch > args.warmup_epochs
        epoch_start = time.perf_counter()
        train_metrics = train_epoch(model, train_loader, optimizer, criterion, args)
        scheduler.step()
        val_metrics = evaluate(model, val_loader, criterion)
        epoch_seconds = time.perf_counter() - epoch_start
        training_seconds += epoch_seconds
        if val_metrics["OA"] >= best_val:
            # checkpoint选择只看验证集OA，测试集在训练结束前完全不可见。
            best_val, best_epoch = val_metrics["OA"], epoch
            payload = checkpoint_payload(
                model, optimizer, scheduler, epoch, best_val, best_epoch, args
            )
            payload["training_seconds"] = training_seconds
            payload["split_sha256"] = split_digest
            atomic_save(payload, best_path)
        row = {
            "epoch": epoch,
            "phase": "warmup" if is_dynamic_method(args.method) and epoch <= args.warmup_epochs else args.method,
            "dynamic_enabled": int(is_dynamic_method(args.method) and epoch > args.warmup_epochs),
            "learning_rate": optimizer.param_groups[0]["lr"],
            **train_metrics,
            "val_loss": val_metrics["loss"],
            "val_OA": val_metrics["OA"],
            "val_mean_max_weight": val_metrics["mean_max_weight"],
            "seconds": epoch_seconds,
        }
        append_epoch_log(epoch_log, row)
        payload = checkpoint_payload(
            model, optimizer, scheduler, epoch, best_val, best_epoch, args
        )
        payload["training_seconds"] = training_seconds
        payload["split_sha256"] = split_digest
        atomic_save(payload, latest_path)
        # 每轮都保存latest，所以正常在一轮打印结束后Ctrl+C可从下一轮继续。
        print(
            f"epoch={epoch:03d} method={args.method} train_OA={train_metrics['train_OA']:.4f} "
            f"val_OA={val_metrics['OA']:.4f} best={best_val:.4f}@{best_epoch} "
            f"time={epoch_seconds:.1f}s",
            flush=True,
        )

    best_state = torch.load(best_path, map_location="cpu", weights_only=False)
    # 训练结束后重新载入验证集最优模型，而不是默认使用第300轮。
    model.load_state_dict(best_state["model_state_dict"], strict=True)
    if is_dynamic_method(args.method):
        model.dynamic_enabled = int(best_state["epoch"]) > args.warmup_epochs
    macs, params = compute_macs_params(model, dataset, args.patch_size)
    training_summary = {
        "status": "training_complete",
        "method": args.method,
        "seed": args.seed,
        "epochs": args.epochs,
        "best_epoch": int(best_state["epoch"]),
        "best_val_OA": float(best_state["best_val_OA"]),
        # Report the complete 300-epoch wall time accumulated in latest.pt,
        # not the partial time stored when the best checkpoint was selected.
        "training_seconds": float(training_seconds),
        "parameters": params,
        "macs": macs,
        "split_sha256": split_digest,
    }
    (method_dir / "training_summary.json").write_text(
        json.dumps(training_summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    if args.skip_full_test:
        # 调参阶段可只训练和验证，避免反复查看测试集造成信息泄漏。
        print(json.dumps(training_summary, ensure_ascii=False, indent=2), flush=True)
        return

    full_test_dir = method_dir / f"full_test_epoch_{best_epoch}"
    model_kind = {
        "baseline": "baseline_mcf3",
        "dynamic": "warmup_dynamic",
        "dynamic_final": "final_fusion_dynamic",
    }[args.method]
    command = [
        sys.executable,
        str(FULL_TEST_PATH),
        "--model-kind", model_kind,
        "--checkpoint", str(best_path),
        "--data-dir", str(args.data_dir),
        "--output-dir", str(full_test_dir),
        "--batch-size", str(args.eval_batch_size),
        "--patch-size", str(args.patch_size),
        "--progress-interval", "100",
    ]
    if args.dataset != "houston2018":
        # SZUTree不是缓存中的固定官方test，因此显式传入本次划分的test_idx。
        command.extend(["--index-file", str(method_dir / "test_idx.npy")])
    print("Running official full test:\n" + subprocess.list2cmdline(command), flush=True)
    subprocess.run(command, check=True, cwd=str(ROOT))
    # 全测试完成后再把指标整理成原MSAF风格TXT和单实验Excel。
    metrics = json.loads((full_test_dir / "metrics.json").read_text(encoding="utf-8"))
    write_msaf_report(
        method_dir / f"{dataset_tag}_3mod_{args.method}_seed{args.seed}_Report.txt",
        train_counts, val_counts, metrics,
        training_summary["training_seconds"], metrics["elapsed_seconds"], params, macs,
    )
    write_excel(
        method_dir / f"{dataset_tag}_3mod_{args.method}_seed{args.seed}.xlsx",
        args, split_info, training_summary, metrics, epoch_log,
    )
    final_summary = {**training_summary, "official_test": {
        key: metrics[key] for key in ("OA", "AA", "kappa", "samples", "elapsed_seconds")
    }}
    (method_dir / "final_summary.json").write_text(
        json.dumps(final_summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(final_summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
